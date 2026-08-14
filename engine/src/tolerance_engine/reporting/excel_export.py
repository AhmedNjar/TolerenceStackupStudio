"""
Excel export via openpyxl.

Three different "liveness" tiers exist in the exported workbook, and they
are NOT the same guarantee — this is worth being precise about rather than
just labeling the whole sheet "live":

1. NOMINAL is fully live for ANY change (tolerance or nominal value) — it's
   a direct expression substitution via excel_formula_printer, so Excel
   recomputes it exactly the way the engine would.

2. WORST CASE and RSS are live with respect to TOLERANCE and DISTRIBUTION
   TYPE changes (exact — those enter the formulas as simple cell
   references), but the partial-derivative COEFFICIENTS baked into those
   formulas are frozen at export time. For a purely linear/additive closing
   equation (e.g. "A - B - C", by far the common case for 1D/multi-gap
   chains) the derivatives are constants anyway, so this is exact
   regardless of nominal changes too. For a NONLINEAR closing equation
   (e.g. the triangular vector loop template, which has a cos() term), the
   WC/RSS formulas stay exact for tolerance edits but become an
   increasingly rough approximation if you change a NOMINAL value far from
   what it was at export time — Excel can't do the symbolic
   re-differentiation that would require. A cell comment on each WC/RSS
   result cell says this explicitly.

3. MONTE CARLO (mean, stdDev, skewness, kurtosis, Cp/Cpk, yield%, DPPM,
   histogram) is a STATIC SNAPSHOT from the simulation run at export time —
   there's no meaningful way to make a random simulation "live" in a
   spreadsheet. Re-run the export to refresh it.
"""
from __future__ import annotations

from dataclasses import dataclass
from typing import Optional

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from tolerance_engine.expression_engine.parser import (
    build_symbols,
    parse_expression,
    partial_derivatives,
)
from tolerance_engine.models import AnalysisMethod, ClosingEquation, Component
from tolerance_engine.reporting.excel_formula_printer import to_excel_formula
from tolerance_engine.statistics.monte_carlo import monte_carlo

_SIGMA_DIVISOR_FORMULA = (
    'IF(G{r}="NORMAL_3S",3,IF(G{r}="NORMAL_6S",6,'
    'IF(G{r}="UNIFORM",SQRT(3),IF(G{r}="TRIANGULAR",SQRT(6),NA()))))'
)

_WC_RSS_LIVENESS_NOTE = (
    "Live for tolerance/distribution edits. For a NONLINEAR closing equation, this becomes "
    "approximate if you change a NOMINAL value far from its value at export time — the "
    "sensitivity coefficients baked into this formula are frozen at export. Re-export to refresh "
    "them. Exact regardless of nominal changes for a purely linear closing equation."
)


@dataclass
class _ComponentRowRefs:
    row: int
    label_cell: str
    nominal_cell: str
    effective_nominal_cell: str
    upper_tol_cell: str
    lower_tol_cell: str
    sigma_cell: str
    mean_shift_cell: str


def build_workbook(
    components: list[Component],
    closing_equations: list[ClosingEquation],
    monte_carlo_runs: int,
    random_seed: Optional[int],
) -> Workbook:
    wb = Workbook()
    components_sheet = wb.active
    components_sheet.title = "Components"
    row_refs = _write_components_sheet(components_sheet, components)

    analysis_sheet = wb.create_sheet("Analysis")
    _write_analysis_sheet(analysis_sheet, components, closing_equations, row_refs, monte_carlo_runs, random_seed, wb)

    return wb


def _write_components_sheet(sheet: Worksheet, components: list[Component]) -> dict[str, _ComponentRowRefs]:
    headers = [
        "Label", "Name", "Kind", "Nominal", "+UpperTol", "-LowerTol", "Distribution",
        "Thermal Alpha", "Thermal DeltaT", "Effective Nominal", "Sigma Divisor", "Sigma", "Mean Shift",
    ]
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)

    row_refs: dict[str, _ComponentRowRefs] = {}
    for i, c in enumerate(components):
        r = i + 2
        sheet.cell(row=r, column=1, value=c.label)
        sheet.cell(row=r, column=2, value=c.name)
        sheet.cell(row=r, column=3, value=c.kind.value)
        sheet.cell(row=r, column=4, value=c.nominal)
        sheet.cell(row=r, column=5, value=c.upperTol)
        sheet.cell(row=r, column=6, value=c.lowerTol)
        sheet.cell(row=r, column=7, value=c.distribution.value)
        sheet.cell(row=r, column=8, value=c.thermal.alpha if c.thermal else None)
        sheet.cell(row=r, column=9, value=c.thermal.deltaT if c.thermal else None)

        if c.thermal:
            eff_nominal_formula = f"=D{r}*(1+H{r}*I{r})"
        else:
            eff_nominal_formula = f"=D{r}"
        sheet.cell(row=r, column=10, value=eff_nominal_formula)

        sheet.cell(row=r, column=11, value="=" + _SIGMA_DIVISOR_FORMULA.format(r=r))
        sheet.cell(row=r, column=12, value=f"=((E{r}+F{r})/2)/K{r}")
        sheet.cell(row=r, column=13, value=f"=(E{r}-F{r})/2")

        row_refs[c.label] = _ComponentRowRefs(
            row=r,
            label_cell=f"'Components'!$A${r}",
            nominal_cell=f"'Components'!$D${r}",
            effective_nominal_cell=f"'Components'!$J${r}",
            upper_tol_cell=f"'Components'!$E${r}",
            lower_tol_cell=f"'Components'!$F${r}",
            sigma_cell=f"'Components'!$L${r}",
            mean_shift_cell=f"'Components'!$M${r}",
        )

    for col in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 16

    return row_refs


def _write_analysis_sheet(
    sheet: Worksheet,
    components: list[Component],
    closing_equations: list[ClosingEquation],
    row_refs: dict[str, _ComponentRowRefs],
    monte_carlo_runs: int,
    random_seed: Optional[int],
    workbook: Workbook,
) -> None:
    header_font = Font(bold=True)
    row = 1

    for closing in closing_equations:
        sheet.cell(row=row, column=1, value=f"Closing dimension: {closing.label} = {closing.expression}").font = header_font
        row += 1

        symbols = build_symbols(components)
        expr = parse_expression(closing.expression, symbols)
        derivatives = partial_derivatives(expr, symbols, components)
        cell_refs = {c.label: row_refs[c.label].effective_nominal_cell for c in components}
        nominal_formula = to_excel_formula(expr, cell_refs)

        sheet.cell(row=row, column=1, value="Nominal (live)")
        nominal_cell_ref = f"B{row}"
        sheet.cell(row=row, column=2, value=nominal_formula)
        row += 1

        wc_min_terms, wc_max_terms = [], []
        rss_variance_terms = []
        rss_mean_shift_terms = []
        contribution_terms: dict[str, str] = {}
        for c in components:
            refs = row_refs[c.label]
            d = derivatives[c.label]
            if d >= 0:
                wc_max_terms.append(f"{d}*{refs.upper_tol_cell}")
                wc_min_terms.append(f"{d}*(-{refs.lower_tol_cell})")
            else:
                wc_max_terms.append(f"{d}*(-{refs.lower_tol_cell})")
                wc_min_terms.append(f"{d}*{refs.upper_tol_cell}")
            rss_variance_terms.append(f"({d}*{refs.sigma_cell})^2")
            rss_mean_shift_terms.append(f"{d}*{refs.mean_shift_cell}")
            contribution_terms[c.id] = f"({d}*{refs.sigma_cell})^2"

        sheet.cell(row=row, column=1, value="Worst Case zMin (live*)")
        c1 = sheet.cell(row=row, column=2, value=f"={nominal_cell_ref}+{'+'.join(wc_min_terms)}")
        c1.comment = Comment(_WC_RSS_LIVENESS_NOTE, "tolerance-engine")
        row += 1

        sheet.cell(row=row, column=1, value="Worst Case zMax (live*)")
        c2 = sheet.cell(row=row, column=2, value=f"={nominal_cell_ref}+{'+'.join(wc_max_terms)}")
        c2.comment = Comment(_WC_RSS_LIVENESS_NOTE, "tolerance-engine")
        row += 1

        sheet.cell(row=row, column=1, value="RSS sigma (live*)")
        sigma_cell = f"B{row}"
        c3 = sheet.cell(row=row, column=2, value=f"=SQRT({'+'.join(rss_variance_terms)})")
        c3.comment = Comment(_WC_RSS_LIVENESS_NOTE, "tolerance-engine")
        row += 1

        sheet.cell(row=row, column=1, value="RSS predicted range (live*, ±3σ incl. mean shift)")
        mean_shift_sum = "+".join(rss_mean_shift_terms) if rss_mean_shift_terms else "0"
        sheet.cell(row=row, column=2, value=f"={nominal_cell_ref}+({mean_shift_sum})-3*{sigma_cell}")
        sheet.cell(row=row, column=3, value=f"={nominal_cell_ref}+({mean_shift_sum})+3*{sigma_cell}")
        row += 1

        row += 1  # blank separator before contribution + Monte Carlo snapshot

        sheet.cell(row=row, column=1, value="Variance contribution % (live*, Pareto)").font = header_font
        row += 1
        total_variance_expr = "+".join(contribution_terms.values()) if contribution_terms else "1"
        for c in components:
            sheet.cell(row=row, column=1, value=c.label)
            sheet.cell(row=row, column=2, value=f"=100*({contribution_terms[c.id]})/({total_variance_expr})")
            row += 1

        row += 1

        mc_result = monte_carlo(closing.expression, components, monte_carlo_runs, closing.specLimits, random_seed)
        sheet.cell(row=row, column=1, value=f"Monte Carlo snapshot ({mc_result['runs']} runs, STATIC)").font = header_font
        row += 1
        for label, key in [
            ("Mean", "mean"), ("Std Dev", "stdDev"), ("Skewness", "skewness"), ("Kurtosis", "kurtosis"),
            ("Cp", "cp"), ("Cpk", "cpk"), ("Yield %", "yieldPct"), ("DPPM", "dppm"),
        ]:
            sheet.cell(row=row, column=1, value=label)
            sheet.cell(row=row, column=2, value=mc_result[key])
            row += 1

        row = _write_histogram_and_chart(workbook, sheet, closing.id, mc_result["histogram"], row)

        row += 2  # gap before the next closing equation

    for col, width in enumerate([40, 18, 18], start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width


def _write_histogram_and_chart(workbook: Workbook, analysis_sheet: Worksheet, closing_id: str, histogram: dict, start_row: int) -> int:
    data_sheet_name = f"hist_{closing_id}"[:31]  # Excel sheet name length limit
    data_sheet = workbook.create_sheet(data_sheet_name)
    data_sheet.sheet_state = "hidden"

    data_sheet.cell(row=1, column=1, value="Bin midpoint")
    data_sheet.cell(row=1, column=2, value="Count")
    bin_edges = histogram["binEdges"]
    counts = histogram["counts"]
    for i, count in enumerate(counts):
        midpoint = (bin_edges[i] + bin_edges[i + 1]) / 2.0
        data_sheet.cell(row=i + 2, column=1, value=midpoint)
        data_sheet.cell(row=i + 2, column=2, value=count)

    chart = BarChart()
    chart.title = f"{closing_id} Monte Carlo distribution"
    chart.x_axis.title = "Value"
    chart.y_axis.title = "Count"
    chart.style = 10
    data_ref = Reference(data_sheet, min_col=2, min_row=1, max_row=len(counts) + 1)
    categories_ref = Reference(data_sheet, min_col=1, min_row=2, max_row=len(counts) + 1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(categories_ref)
    chart.width, chart.height = 18, 8

    analysis_sheet.add_chart(chart, f"E{start_row}")
    return start_row + 18  # approx chart height in rows, so the next closing equation doesn't overlap
